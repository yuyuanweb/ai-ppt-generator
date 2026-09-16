import base64
import json
import struct
import uuid
import zlib
from collections.abc import AsyncGenerator

import httpx
import pytest
from httpx import ASGITransport, AsyncClient
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.core.config import get_settings
from app.core.db import async_session_factory
from app.domain.content import Deck, ImageBlock, Slide, TextBlock
from app.images.bailian import BailianImageProvider, _closest_size, _resolve_base_url
from app.images.base import ImageAsset, ImageRequest
from app.images.generated import GeneratedImageProvider
from app.images.pipeline import ImagePipeline
from app.images.unsplash import UnsplashImageProvider, sanitize_unsplash_query
from app.images.validate import ImageRejected, validate_image
from app.main import app
from app.models.project import Project, ProjectOutline
from app.models.slide import Slide as SlideRow
from app.render.pptx import render_deck_to_pptx
from app.services.media import MEDIA_PREFIX, media_url, store_image
from app.services.slide_images import resolve_slide_images
from app.storage import get_storage


def _chunk(tag: bytes, data: bytes) -> bytes:
    crc = struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)
    return struct.pack(">I", len(data)) + tag + data + crc


def minimal_png(*, width: int = 2, height: int = 2) -> bytes:
    """手工构造最小合法 PNG，避免测试依赖外部文件。"""
    signature = b"\x89PNG\r\n\x1a\n"
    ihdr = _chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
    raw = b"".join(b"\x00" + b"\xff\x00\x00" * width for _ in range(height))
    idat = _chunk(b"IDAT", zlib.compress(raw))
    iend = _chunk(b"IEND", b"")
    return signature + ihdr + idat + iend


MINIMAL_PNG = minimal_png()
MINIMAL_JPEG = b"\xff\xd8\xff\xe0" + b"\x00" * 32
MINIMAL_WEBP = b"RIFF" + struct.pack("<I", 12) + b"WEBP" + b"\x00" * 4


@pytest.fixture
async def client() -> AsyncGenerator[AsyncClient, None]:
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as value:
        yield value


async def _sign_up(client: AsyncClient) -> dict[str, str]:
    response = await client.post(
        "/api/v1/auth/register",
        json={"email": f"img_{uuid.uuid4().hex}@example.com", "password": "password123"},
    )
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


async def _project_with_image_slide(
    client: AsyncClient,
    headers: dict[str, str],
    *,
    locked: bool = False,
    revision: int = 1,
    status: str = "ready",
) -> tuple[dict, SlideRow]:
    response = await client.post(
        "/api/v1/projects",
        json={"title": "配图项目", "page_count": 5},
        headers=headers,
    )
    project = response.json()
    page_id = str(uuid.uuid4())
    block_id = "img-1"

    async with async_session_factory() as session:
        result = await session.execute(
            select(Project)
            .options(selectinload(Project.outline))
            .where(Project.id == uuid.UUID(project["id"]))
        )
        record = result.scalar_one()
        session.add(
            ProjectOutline(
                project_id=record.id,
                status="confirmed",
                pages=[
                    {
                        "id": page_id,
                        "title": "示意图页",
                        "objective": "展示配图",
                        "key_points": ["要点"],
                        "source_refs": [],
                        "layout_id": "image-right",
                    }
                ],
                revision=2,
            )
        )
        slide = SlideRow(
            project_id=record.id,
            outline_page_id=uuid.UUID(page_id),
            position=1,
            layout_id="image-right",
            layout_mode="fixed",
            title="示意图页",
            status=status,
            blocks=[
                {
                    "id": "t1",
                    "slot_id": "title",
                    "type": "text",
                    "text": "标题",
                    "locked": False,
                },
                {
                    "id": block_id,
                    "slot_id": "image",
                    "type": "image",
                    "alt": "业务流程图",
                    "source": "placeholder",
                    "url": None,
                    "credit": None,
                    "locked": locked,
                },
            ],
            issues=[],
            revision=revision,
        )
        session.add(slide)
        record.status = "ready" if status == "ready" else "generating"
        await session.commit()
        await session.refresh(slide)
        return project, slide


# --- validate_image ---


def test_validate_image_accepts_png_jpeg_webp() -> None:
    assert validate_image(MINIMAL_PNG) == (".png", "image/png")
    assert validate_image(MINIMAL_JPEG) == (".jpg", "image/jpeg")
    assert validate_image(MINIMAL_WEBP) == (".webp", "image/webp")


def test_validate_image_rejects_text_and_oversized(monkeypatch: pytest.MonkeyPatch) -> None:
    with pytest.raises(ImageRejected, match="不支持"):
        validate_image(b"plain text bytes")

    settings = get_settings()
    monkeypatch.setattr(settings, "max_image_mb", 0)
    with pytest.raises(ImageRejected, match="上限"):
        validate_image(MINIMAL_PNG)


# --- ImagePipeline ---


class _FakeProvider:
    def __init__(
        self,
        *,
        source: str,
        available: bool = True,
        asset: ImageAsset | None = None,
    ) -> None:
        self.source = source  # type: ignore[assignment]
        self._available = available
        self._asset = asset
        self.calls = 0

    def available(self) -> bool:
        return self._available

    async def fetch(self, request: ImageRequest) -> ImageAsset | None:
        self.calls += 1
        return self._asset


@pytest.mark.asyncio
async def test_pipeline_prefers_generated_then_stock() -> None:
    generated = _FakeProvider(
        source="generated",
        asset=ImageAsset(data=MINIMAL_PNG, content_type="image/png", source="generated"),
    )
    stock = _FakeProvider(
        source="stock",
        asset=ImageAsset(data=MINIMAL_JPEG, content_type="image/jpeg", source="stock"),
    )
    pipeline = ImagePipeline([generated, stock])  # type: ignore[arg-type]
    asset = await pipeline.fetch(ImageRequest(prompt="p", query="q", aspect_ratio=1.5))
    assert asset is not None
    assert asset.source == "generated"
    assert stock.calls == 0


@pytest.mark.asyncio
async def test_pipeline_falls_back_when_generated_unavailable_or_none() -> None:
    stock_asset = ImageAsset(
        data=MINIMAL_JPEG,
        content_type="image/jpeg",
        source="stock",
        credit="Photo by Ada on Unsplash",
    )
    unavailable = _FakeProvider(source="generated", available=False)
    returns_none = _FakeProvider(source="generated", asset=None)
    stock = _FakeProvider(source="stock", asset=stock_asset)

    via_unavailable = await ImagePipeline([unavailable, stock]).fetch(  # type: ignore[arg-type]
        ImageRequest(prompt="p", query="q", aspect_ratio=1.0)
    )
    assert via_unavailable is not None and via_unavailable.source == "stock"

    stock.calls = 0
    via_none = await ImagePipeline([returns_none, stock]).fetch(  # type: ignore[arg-type]
        ImageRequest(prompt="p", query="q", aspect_ratio=1.0)
    )
    assert via_none is not None and via_none.source == "stock"


@pytest.mark.asyncio
async def test_pipeline_returns_none_when_all_unavailable() -> None:
    pipeline = ImagePipeline(
        [
            _FakeProvider(source="generated", available=False),  # type: ignore[list-item]
            _FakeProvider(source="stock", available=False),  # type: ignore[list-item]
        ]
    )
    assert await pipeline.fetch(ImageRequest(prompt="p", query="q", aspect_ratio=1.0)) is None


# --- GeneratedImageProvider ---


@pytest.mark.asyncio
async def test_generated_provider_reads_b64_and_url() -> None:
    png = MINIMAL_PNG
    generation_calls = 0

    def handler(request: httpx.Request) -> httpx.Response:
        nonlocal generation_calls
        if request.url.path.endswith("/images/generations"):
            generation_calls += 1
            if generation_calls == 1:
                return httpx.Response(
                    200,
                    json={"data": [{"b64_json": base64.b64encode(png).decode()}]},
                )
            return httpx.Response(
                200,
                json={"data": [{"url": "https://cdn.example.com/gen.png"}]},
            )
        if "cdn.example.com" in str(request.url):
            return httpx.Response(200, content=png)
        return httpx.Response(404)

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        provider = GeneratedImageProvider(
            client=client,
            base_url="https://api.example.com/v1",
            api_key="test-key",
            model="gpt-image-1",
        )
        first = await provider.fetch(ImageRequest(prompt="a", query="a", aspect_ratio=1.5))
        second = await provider.fetch(ImageRequest(prompt="b", query="b", aspect_ratio=1.0))

    assert first is not None and first.data == png and first.source == "generated"
    assert second is not None and second.data == png


@pytest.mark.asyncio
async def test_generated_provider_http_500_returns_none() -> None:
    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(500, text="boom")

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        provider = GeneratedImageProvider(
            client=client,
            base_url="https://api.example.com/v1",
            api_key="test-key",
            model="gpt-image-1",
        )
        result = await provider.fetch(ImageRequest(prompt="a", query="a", aspect_ratio=1.0))
    assert result is None


# --- BailianImageProvider ---


def test_bailian_size_and_workspace_base_url() -> None:
    assert _closest_size(1.0) == "1328*1328"
    assert _closest_size(1.5) == "1664*928"
    assert _closest_size(0.67) == "928*1664"
    assert (
        _resolve_base_url("https://dashscope.aliyuncs.com/api/v1", "ws-demo")
        == "https://ws-demo.cn-beijing.maas.aliyuncs.com/api/v1"
    )


@pytest.mark.asyncio
async def test_bailian_provider_posts_native_payload_and_downloads_image() -> None:
    png = MINIMAL_PNG
    captured: dict[str, object] = {}

    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.path.endswith("/services/aigc/multimodal-generation/generation"):
            captured["auth"] = request.headers.get("Authorization")
            captured["json"] = json.loads(request.content.decode())
            return httpx.Response(
                200,
                json={
                    "output": {
                        "choices": [
                            {
                                "finish_reason": "stop",
                                "message": {
                                    "role": "assistant",
                                    "content": [{"image": "https://cdn.example.com/bailian.png"}],
                                },
                            }
                        ]
                    },
                    "request_id": "req-1",
                },
            )
        if "cdn.example.com" in str(request.url):
            return httpx.Response(200, content=png)
        return httpx.Response(404)

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        provider = BailianImageProvider(
            client=client,
            api_key="sk-test",
            model="qwen-image-3.0",
            base_url="https://dashscope.aliyuncs.com/api/v1",
        )
        asset = await provider.fetch(
            ImageRequest(prompt="商务配图，无文字", query="business", aspect_ratio=1.5)
        )

    assert asset is not None and asset.data == png and asset.source == "generated"
    assert captured["auth"] == "Bearer sk-test"
    body = captured["json"]
    assert isinstance(body, dict)
    assert body["model"] == "qwen-image-3.0"
    assert body["input"]["messages"][0]["content"][0]["text"] == "商务配图，无文字"
    assert body["parameters"]["size"] == "1664*928"
    assert body["parameters"]["n"] == 1
    assert body["parameters"]["watermark"] is False
    assert body["parameters"]["prompt_extend"] is True


@pytest.mark.asyncio
async def test_bailian_provider_business_error_returns_none() -> None:
    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={"code": "InvalidApiKey", "message": "Invalid API-key provided."},
        )

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        provider = BailianImageProvider(
            client=client,
            api_key="bad",
            model="qwen-image-3.0",
        )
        result = await provider.fetch(ImageRequest(prompt="a", query="a", aspect_ratio=1.0))
    assert result is None


@pytest.mark.asyncio
async def test_bailian_provider_http_500_returns_none() -> None:
    def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(500, text="boom")

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        provider = BailianImageProvider(
            client=client,
            api_key="sk-test",
            model="qwen-image-3.0",
        )
        result = await provider.fetch(ImageRequest(prompt="a", query="a", aspect_ratio=1.0))
    assert result is None


# --- UnsplashImageProvider ---


@pytest.mark.asyncio
async def test_unsplash_parses_image_credit_and_tracks_download() -> None:
    tracked: list[str] = []
    png = MINIMAL_PNG

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "/search/photos" in url:
            return httpx.Response(
                200,
                json={
                    "results": [
                        {
                            "urls": {"regular": "https://images.unsplash.com/photo.jpg"},
                            "links": {
                                "download_location": "https://api.unsplash.com/photos/x/download"
                            },
                            "user": {"name": "Ada Lovelace"},
                        }
                    ]
                },
            )
        if "images.unsplash.com" in url:
            return httpx.Response(200, content=png, headers={"content-type": "image/png"})
        if "/download" in url:
            tracked.append(url)
            return httpx.Response(200, json={"url": "https://images.unsplash.com/photo.jpg"})
        return httpx.Response(404)

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        provider = UnsplashImageProvider(client=client, access_key="unsplash-key")
        asset = await provider.fetch(ImageRequest(prompt="p", query="workflow", aspect_ratio=1.6))

    assert asset is not None
    assert asset.data == png
    assert asset.source == "stock"
    assert asset.credit == "Photo by Ada Lovelace on Unsplash"
    assert tracked and "download" in tracked[0]


@pytest.mark.asyncio
async def test_unsplash_empty_results_returns_none() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, json={"results": []})

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        provider = UnsplashImageProvider(client=client, access_key="unsplash-key")
        assert await provider.fetch(ImageRequest(prompt="p", query="q", aspect_ratio=1.0)) is None


def test_sanitize_unsplash_query_strips_fullwidth_colon() -> None:
    cleaned = sanitize_unsplash_query(
        "抽象几何图形：三条交错的路径分别代表认知、行为与环境，寓意三者协同驱动成长"
    )
    assert "：" not in cleaned
    assert ":" not in cleaned
    assert "抽象几何图形" in cleaned
    assert " " in cleaned


@pytest.mark.asyncio
async def test_unsplash_retries_after_content_removed() -> None:
    """首个 query 若仍被 410，应继续尝试短候选。"""
    seen: list[str] = []
    png = MINIMAL_PNG

    def handler(request: httpx.Request) -> httpx.Response:
        url = str(request.url)
        if "/search/photos" in url:
            query = request.url.params.get("query", "")
            seen.append(query)
            if len(seen) == 1:
                return httpx.Response(410, json={"errors": ["Content removed"]})
            return httpx.Response(
                200,
                json={
                    "results": [
                        {
                            "urls": {"regular": "https://images.unsplash.com/photo.jpg"},
                            "links": {},
                            "user": {"name": "Ada"},
                        }
                    ]
                },
            )
        if "images.unsplash.com" in url:
            return httpx.Response(200, content=png, headers={"content-type": "image/png"})
        return httpx.Response(404)

    transport = httpx.MockTransport(handler)
    async with httpx.AsyncClient(transport=transport) as client:
        provider = UnsplashImageProvider(client=client, access_key="unsplash-key")
        asset = await provider.fetch(
            ImageRequest(
                prompt="p",
                query="抽象几何图形：三条交错的路径分别代表认知、行为与环境",
                aspect_ratio=1.6,
            )
        )

    assert asset is not None
    assert asset.source == "stock"
    assert len(seen) >= 2
    assert all("：" not in q for q in seen)


# --- resolve_slide_images ---


@pytest.mark.asyncio
async def test_resolve_slide_images_writes_url_source_credit() -> None:
    user_id = uuid.uuid4()
    project_id = uuid.uuid4()
    provider = _FakeProvider(
        source="stock",
        asset=ImageAsset(
            data=MINIMAL_PNG,
            content_type="image/png",
            source="stock",
            credit="Photo by Ada on Unsplash",
        ),
    )
    slide = Slide(
        id=str(uuid.uuid4()),
        layout_id="image-right",
        blocks=[
            TextBlock(id="t1", slot_id="title", text="标题"),
            ImageBlock(id="i1", slot_id="image", alt="流程图", source="placeholder"),
        ],
    )
    resolved = await resolve_slide_images(
        ImagePipeline([provider]),  # type: ignore[arg-type]
        user_id=user_id,
        project_id=project_id,
        deck_title="季度复盘",
        page_title="示意图页",
        slide=slide,
    )
    image = next(block for block in resolved.blocks if block.type == "image")
    assert image.url is not None and image.url.startswith("/api/v1/media/media/")
    assert image.source == "stock"
    assert image.credit == "Photo by Ada on Unsplash"
    # 原对象不得被就地修改
    assert slide.blocks[1].url is None


@pytest.mark.asyncio
async def test_resolve_slide_images_keeps_placeholder_and_skips_locked() -> None:
    user_id = uuid.uuid4()
    project_id = uuid.uuid4()
    empty = _FakeProvider(source="generated", asset=None)
    slide = Slide(
        id=str(uuid.uuid4()),
        layout_id="image-right",
        blocks=[
            ImageBlock(id="i1", slot_id="image", alt="a", source="placeholder"),
            ImageBlock(
                id="i2",
                slot_id="image",
                alt="b",
                source="upload",
                url="/api/v1/media/media/x.png",
                locked=True,
            ),
        ],
    )
    resolved = await resolve_slide_images(
        ImagePipeline([empty]),  # type: ignore[arg-type]
        user_id=user_id,
        project_id=project_id,
        deck_title="t",
        page_title="p",
        slide=slide,
    )
    assert resolved.blocks[0].url is None
    assert resolved.blocks[0].source == "placeholder"
    assert resolved.blocks[1].url == "/api/v1/media/media/x.png"
    assert empty.calls == 1


# --- media endpoint ---


@pytest.mark.asyncio
async def test_media_endpoint_serves_media_prefix_only(client: AsyncClient) -> None:
    key = store_image(
        user_id=uuid.uuid4(),
        project_id=uuid.uuid4(),
        data=MINIMAL_PNG,
        extension=".png",
    )
    ok = await client.get(f"/api/v1/media/{key}")
    assert ok.status_code == 200
    assert ok.content == MINIMAL_PNG
    assert ok.headers["cache-control"] == "public, max-age=31536000, immutable"

    upload_key = f"uploads/{uuid.uuid4().hex}.png"
    get_storage().save(upload_key, MINIMAL_PNG)
    denied = await client.get(f"/api/v1/media/{upload_key}")
    assert denied.status_code == 404

    missing = await client.get(f"/api/v1/media/{MEDIA_PREFIX}missing/{uuid.uuid4().hex}.png")
    assert missing.status_code == 404


# --- replace image endpoint ---


@pytest.mark.asyncio
async def test_replace_image_success_and_conflicts(client: AsyncClient) -> None:
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers)
    block_id = "img-1"

    ok = await client.put(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/{block_id}/image",
        headers=headers,
        data={"revision": str(slide.revision)},
        files={"file": ("shot.png", MINIMAL_PNG, "image/png")},
    )
    assert ok.status_code == 200
    body = ok.json()
    assert body["revision"] == slide.revision + 1
    image = next(block for block in body["blocks"] if block["id"] == block_id)
    assert image["source"] == "upload"
    assert image["locked"] is True
    assert image["url"].startswith("/api/v1/media/media/")
    assert image["credit"] is None

    conflict = await client.put(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/{block_id}/image",
        headers=headers,
        data={"revision": str(slide.revision)},
        files={"file": ("shot.png", MINIMAL_PNG, "image/png")},
    )
    assert conflict.status_code == 409

    bad = await client.put(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/{block_id}/image",
        headers=headers,
        data={"revision": str(body["revision"])},
        files={"file": ("note.txt", b"not an image", "text/plain")},
    )
    assert bad.status_code == 422
    assert "不支持" in bad.json()["detail"]


@pytest.mark.asyncio
async def test_replace_image_rejects_while_generating(client: AsyncClient) -> None:
    headers = await _sign_up(client)
    project, slide = await _project_with_image_slide(client, headers, status="generating")

    response = await client.put(
        f"/api/v1/projects/{project['id']}/deck/slides/{slide.id}/blocks/img-1/image",
        headers=headers,
        data={"revision": str(slide.revision)},
        files={"file": ("shot.png", MINIMAL_PNG, "image/png")},
    )
    assert response.status_code == 409
    assert response.json()["detail"] == "页面正在生成中，请稍后再换图"


# --- PPTX with real image ---


def test_pptx_export_embeds_real_picture() -> None:
    user_id = uuid.uuid4()
    project_id = uuid.uuid4()
    key = store_image(
        user_id=user_id,
        project_id=project_id,
        data=minimal_png(width=8, height=4),
        extension=".png",
    )
    deck = Deck(
        id="d1",
        title="带图文稿",
        theme_id="midnight",
        slides=[
            Slide(
                id="s1",
                layout_id="image-right",
                blocks=[
                    TextBlock(id="t1", slot_id="title", text="标题"),
                    ImageBlock(
                        id="i1",
                        slot_id="image",
                        alt="示意图",
                        source="upload",
                        url=media_url(key),
                    ),
                ],
            )
        ],
    )
    presentation = Presentation(render_deck_to_pptx(deck))
    pictures = [
        shape
        for shape in presentation.slides[0].shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert pictures
    # 宽图应对左右裁切
    picture = pictures[0]
    assert picture.crop_left > 0
    assert picture.crop_right > 0
